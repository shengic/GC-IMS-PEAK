# -*- coding: utf-8 -*-
"""把一份「內容 DSL」渲染成 Word 檔。

兩個 Word 專屬的坑，都不會報錯、只會默默不生效：

1. **中文字型要設 `w:eastAsia`。** `run.font.name` 只寫 ascii/hAnsi 兩個屬性，
   Word 遇到中文字會改用「東亞字型」那一欄——沒設的話整份中文都會退回新細明體，
   而您指定的 Iansui 只套用在英數字上。

2. **OMML 必須掛在 `m:` 命名空間下。** 命名空間宣告錯了，Word 會當成未知標籤
   整段吞掉，方程式直接不見，也沒有任何錯誤訊息。

Version: 1.0 — by Albert Sheng（手冊建置，2026-09-21）
"""
import copy

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK,
                            WD_TAB_ALIGNMENT)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FONT = "Iansui"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT2 = RGBColor(0x0B, 0x6E, 0x4F)
WARNC = RGBColor(0xA5, 0x40, 0x1A)

BAND = "F3F5F8"
WARNBG = "FDF3EC"
OKBG = "EEF6F2"
HEADBG = "EAEFF5"


# --------------------------------------------------------------------------- #
# 字型 / 底色 / 框線
# --------------------------------------------------------------------------- #
def set_font(run, size=10.5, colour=INK, bold=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.font.bold = bold
    # **關鍵**：東亞字型要另外指定，否則中文不會套用
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), name)


def shade(cell_or_par, hexcolour):
    el = cell_or_par._element
    pr = el.find(qn("w:tcPr")) if el.tag.endswith("tc") else None
    if pr is None:
        pr = el.get_or_add_tcPr() if el.tag.endswith("tc") \
            else el.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolour)
    pr.append(sh)


def left_bar(par, hexcolour):
    """段落左側的彩色直條（用段落框線做）。"""
    pr = par._element.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hexcolour)
    bd.append(left)
    pr.append(bd)


# --------------------------------------------------------------------------- #
# 粗體標記解析：文字裡的 **這樣** 會變成強調（用顏色，Iansui 沒有粗體字重）
# --------------------------------------------------------------------------- #
def rich(par, text, size=10.5, colour=INK, emph=ACCENT):
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        r = par.add_run(chunk)
        # 奇數段＝被 ** 包住的部分
        set_font(r, size, emph if i % 2 else colour, bold=bool(i % 2))


# --------------------------------------------------------------------------- #
# OMML 方程式
# --------------------------------------------------------------------------- #
def _m(tag, *children, **attrs):
    el = OxmlElement("m:%s" % tag)
    for k, v in attrs.items():
        el.set(qn("m:%s" % k), v)
    for c in children:
        el.append(c)
    return el


def mrun(text):
    """一段普通的數學文字。"""
    rpr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")          # plain，不要自動斜體
    rpr.append(sty)
    t = OxmlElement("m:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    return _m("r", rpr, t)


def mfrac(num_parts, den_parts):
    """分數。num_parts / den_parts 各是一串 m 元素。"""
    num = _m("num", *num_parts)
    den = _m("den", *den_parts)
    return _m("f", num, den)


def msub(base_parts, sub_parts):
    """下標。"""
    return _m("sSub", _m("e", *base_parts), _m("sub", *sub_parts))


def equation(doc, parts, caption=None):
    """插入一個置中的方程式段落；`parts` 是 m 元素串。"""
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(4)
    omath = _m("oMath", *parts)
    par._p.append(omath)
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        rich(cap, caption, 9, GREY, GREY)
    return par


# --------------------------------------------------------------------------- #
# DSL 渲染
# --------------------------------------------------------------------------- #
def _callout(doc, title, lines, bg, bar_hex, titlecolour, brk=None):
    par = doc.add_paragraph()
    if brk:
        brk.apply(par)
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.left_indent = Cm(0.3)
    left_bar(par, bar_hex)
    shade(par, bg)
    if title:
        r = par.add_run(title)
        set_font(r, 10.5, titlecolour, bold=True)
        par.add_run("\n")
    for i, ln in enumerate(lines):
        if i:
            par.add_run("\n")
        rich(par, ln, 9.5, GREY, titlecolour)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(2)
    gap.add_run("").font.size = Pt(1)


class _Breaker:
    """分頁狀態。

    **不要用「內含分頁符的空白段落」來換頁**——那個空白段落本身會落在新頁的
    最上面，跟表格後面的間隔段落疊在一起時，就會產生整頁只有空白段落的「空白頁」。
    改成把 `page_break_before` 設在**下一個真正有內容的段落**上，
    換頁效果一樣，但不會多出任何段落。
    """

    def __init__(self):
        self.pending = False

    def arm(self):
        self.pending = True

    def apply(self, par):
        if self.pending:
            par.paragraph_format.page_break_before = True
            self.pending = False
        return par


def render(doc, blocks, brk=None):
    brk = brk or _Breaker()

    def newpar(style=None):
        return brk.apply(doc.add_paragraph(style=style) if style
                         else doc.add_paragraph())

    for kind, payload in blocks:
        if kind == "pagebreak":
            brk.arm()

        elif kind == "part":
            num, title, sub = payload
            for i in range(4):
                newpar() if i == 0 else doc.add_paragraph()
            p = doc.add_paragraph()
            set_font(p.add_run("第 %s 部" % num), 13, GREY)
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(10)
            set_font(p2.add_run(title), 30, ACCENT, bold=True)
            p3 = doc.add_paragraph()
            rich(p3, sub, 11, GREY, GREY)
            brk.arm()

        elif kind in ("h2", "h3", "h4"):
            size = {"h2": 18, "h3": 13.5, "h4": 11.5}[kind]
            col = {"h2": ACCENT, "h3": ACCENT2, "h4": INK}[kind]
            p = newpar()
            p.paragraph_format.space_before = Pt(16 if kind == "h2" else 10)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            set_font(p.add_run(payload), size, col, bold=True)

        elif kind == "p":
            p = newpar()
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.35
            rich(p, payload)

        elif kind == "bullets":
            for it in payload:
                p = newpar("List Bullet")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.25
                rich(p, it)

        elif kind == "steps":
            for it in payload:
                p = newpar("List Number")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.25
                rich(p, it)

        elif kind == "note":
            _callout(doc, payload[0], payload[1], BAND, "1F4E79", ACCENT, brk)
        elif kind == "warn":
            _callout(doc, payload[0], payload[1], WARNBG, "A5401A", WARNC, brk)
        elif kind == "ok":
            _callout(doc, payload[0], payload[1], OKBG, "0B6E4F", ACCENT2, brk)

        elif kind == "table":
            header, rows, widths = payload
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(header):
                c = t.rows[0].cells[i]
                c.text = ""
                shade(c, HEADBG)
                rich(c.paragraphs[0], h, 9.5, ACCENT, ACCENT)
            for row in rows:
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = ""
                    rich(cells[i].paragraphs[0], str(val), 9.5)
            for i, w in enumerate(widths):
                for row in t.rows:
                    row.cells[i].width = Cm(w)
            # 表格之間需要一個段落隔開（Word 的限制），但把它縮到 1pt，
            # 否則它自己就可能被擠到下一頁，變成一頁只有一個看不見的空行。
            gap = doc.add_paragraph()
            gap.paragraph_format.space_after = Pt(2)
            gap.paragraph_format.line_spacing = Pt(2)
            for r in gap.runs or [gap.add_run("")]:
                r.font.size = Pt(1)

        elif kind == "eq":
            parts, caption = payload
            equation(doc, parts, caption)

        else:
            raise ValueError("unknown block: %r" % kind)


def new_document(title, company):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    # 預設樣式也要指定東亞字型，否則清單、表格的預設文字會跑掉
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)
    doc.core_properties.title = title
    doc.core_properties.author = "Albert Sheng"
    doc.core_properties.company = company

    _build_footer(sec, title, company)
    return doc


def _page_field(par, size=8.5, colour=GREY):
    """插入 Word 的 PAGE 欄位（會自動顯示當頁頁碼）。"""
    run = par.add_run()
    set_font(run, size, colour)
    for kind, val in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " %s " % val
        run._element.append(el)


def _build_footer(sec, title, company):
    """每一頁的頁尾：公司名（左）｜文件名（中）｜頁碼（右）。

    用**定位點**而不是三個欄位的表格：頁尾放表格在某些 Word 版本會把
    頁面下緣的可用高度吃掉，內文就會莫名其妙提早換頁。
    """
    par = sec.footer.paragraphs[0]
    par.text = ""
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    usable = sec.page_width - sec.left_margin - sec.right_margin
    tabs = par.paragraph_format.tab_stops
    tabs.add_tab_stop(int(usable / 2), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(int(usable), WD_TAB_ALIGNMENT.RIGHT)

    left = par.add_run(company)
    set_font(left, 8.5, ACCENT)
    mid = par.add_run("\t" + title + "\t")
    set_font(mid, 8.5, GREY)
    _page_field(par)

    # 頁尾上方一條細線，和內文分開
    pr = par._element.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), "C8CDD4")
    bd.append(top)
    pr.append(bd)
